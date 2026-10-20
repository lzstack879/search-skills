'''Offline tests: no network, no API key, no real ledger or database.'''

import asyncio
import json
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest

from search_skills import client as client_module
from search_skills.client import (
    DeepSeekClient,
    DeepSeekError,
    FetchPage,
    SearchHit,
    record_usage,
    usage_summary,
)
from search_skills.pipeline import (
    ICP,
    Pipeline,
    extract_emails,
    heuristic_extract,
    heuristic_judge,
    heuristic_queries,
    is_directory,
    parse_json,
    rank_label,
)
from search_skills.report import (
    build_report,
    day_bounds,
    maybe_generate_daily,
    maybe_generate_weekly,
    report_path,
    week_bounds,
)
from search_skills.store import Store, dedupe_key


# ── client ──


def test_search_hit_parses_and_renders():
    hit = SearchHit.from_json({
        'title': 'Auxia',
        'url': 'https://auxia.io',
        'summary': 'AI retention analytics',
        'rerank_score': 2.4,
        'authority': 3,
        'site_name': 'Auxia',
    })
    assert hit.rerank_score == 2.4
    assert hit.authority == 3
    card = hit.card()
    assert 'https://auxia.io' in card
    assert 'AI retention analytics' in card


def test_missing_key_is_a_clear_error():
    with pytest.raises(DeepSeekError):
        asyncio.run(DeepSeekClient(api_key='').search('x'))


def test_search_builds_the_documented_payload(monkeypatch):
    seen = {}

    async def fake_post(url, payload):
        seen['url'] = url
        seen['payload'] = payload
        return {'results': [
            {'title': 'A', 'url': 'https://a.com', 'rerank_score': 1.5, 'authority': 2},
            'not a dict',
        ]}

    client = DeepSeekClient(api_key='k')
    monkeypatch.setattr(client, '_post', fake_post)
    hits = asyncio.run(client.search('site:a.com', 5, need_content=True, ori_question='q'))
    assert seen['url'].endswith('/v1/tools/search')
    assert seen['payload'] == {
        'query': 'site:a.com', 'count': 5, 'need_content': True, 'ori_question': 'q',
    }
    assert len(hits) == 1


def test_search_clamps_count_and_rejects_overlong_queries(monkeypatch):
    captured = {}

    async def fake_post(url, payload):
        captured.update(payload)
        return {'results': []}

    client = DeepSeekClient(api_key='k')
    monkeypatch.setattr(client, '_post', fake_post)
    asyncio.run(client.search('q', 999))
    assert captured['count'] == 50
    with pytest.raises(DeepSeekError):
        asyncio.run(client.search('x' * 129))


def test_fetch_returns_markdown(monkeypatch):
    async def fake_post(url, payload):
        assert payload['url'] == 'https://auxia.io'
        return {'title': 'Auxia', 'content': '# Auxia\n\nWe do retention.'}

    client = DeepSeekClient(api_key='k')
    monkeypatch.setattr(client, '_post', fake_post)
    page = asyncio.run(client.fetch('https://auxia.io'))
    assert page.title == 'Auxia'
    assert 'retention' in page.content
    assert page.length > 0


def test_fetch_rejects_non_http():
    client = DeepSeekClient(api_key='k')
    with pytest.raises(DeepSeekError):
        asyncio.run(client.fetch('not a url'))


def test_usage_ledger_counts_calls(monkeypatch, tmp_path):
    monkeypatch.setattr(client_module, 'USAGE_PATH', tmp_path / 'usage.json')
    record_usage('search')
    record_usage('search')
    record_usage('fetch')
    summary = usage_summary()
    assert summary['today_search'] == 2
    assert summary['today_fetch'] == 1
    assert summary['total_cny'] == pytest.approx(0.015)


# ── store ──


def test_dedupe_key_prefers_website():
    assert dedupe_key('Auxia', 'https://auxia.io') == 'web:https://auxia.io'
    assert dedupe_key('Auxia', '') == 'name:auxia'
    assert dedupe_key('', '', 'A@B.com') == 'email:a@b.com'


def test_store_inserts_then_updates(tmp_path):
    store = Store(tmp_path / 'c.db').init()
    record = {
        'company_name': 'Auxia',
        'company_website': 'https://auxia.io',
        'company_perple': 'Rich Rao（Head of Growth）',
        'company_email': 'rich@auxia.io',
        'industry': 'SaaS',
        'summary': 'retention analytics',
        'match_result': '高（85/100）',
        'match_score': 85,
    }
    assert store.upsert(record) == 'inserted'
    assert store.upsert(record) == 'updated'
    assert store.stats()['companies'] == 1


def test_store_skips_unnamed_rows(tmp_path):
    store = Store(tmp_path / 'c.db').init()
    assert store.upsert({'company_name': '  '}) == 'skipped'


def test_store_lists_only_the_window(tmp_path):
    store = Store(tmp_path / 'c.db').init()
    store.upsert({'company_name': 'A', 'company_website': 'https://a.com'})
    rows = store.list_between('2000-01-01T00:00:00', '2000-01-08T00:00:00')
    assert rows == []
    rows = store.list_between('2000-01-01T00:00:00', '2999-01-01T00:00:00')
    assert len(rows) == 1


# ── report ──


def test_day_bounds_is_one_whole_day():
    from datetime import datetime

    start, end = day_bounds(date(2026, 10, 21))
    begin, finish = datetime.fromisoformat(start), datetime.fromisoformat(end)
    assert finish - begin == timedelta(days=1)
    assert begin.date().toordinal() in (date(2026, 10, 20).toordinal(),
                                        date(2026, 10, 21).toordinal())


def test_week_bounds_is_still_available():
    monday, sunday = week_bounds(date(2026, 10, 21))
    assert monday == date(2026, 10, 19)
    assert sunday == date(2026, 10, 25)


def test_report_path_uses_year_month_day(tmp_path):
    path = report_path(date(2026, 10, 19), tmp_path)
    assert path == tmp_path / '2026' / '10' / '19潜在客户搜索报表.xlsx'


def test_build_report_writes_the_columns(tmp_path):
    from openpyxl import load_workbook

    store = Store(tmp_path / 'c.db').init()
    store.upsert({
        'company_name': 'Auxia',
        'company_website': 'https://auxia.io',
        'company_perple': 'Rich Rao（Head of Growth）',
        'company_email': 'rich@auxia.io',
        'industry': 'SaaS',
        'summary': 'retention analytics',
        'match_result': '高（85/100）',
        'match_score': 85,
    })
    path, rows = build_report(store, date.today(), tmp_path)
    assert rows == 1
    sheet = load_workbook(path).active
    headers = [c.value for c in sheet[1]]
    assert len(headers) == 10
    assert 'company_name' in headers[0]
    assert 'match_result' in headers[6]
    assert 'email_kind' in headers[7]
    assert 'email_verdict' in headers[8]
    assert 'email_grade' in headers[9]
    assert sheet.cell(row=2, column=1).value == 'Auxia'
    assert sheet.cell(row=2, column=7).value == '高（85/100）'


def test_daily_report_is_generated_once(tmp_path):
    store = Store(tmp_path / 'c.db').init()
    first, _rows, created = maybe_generate_daily(store, date(2026, 10, 21), tmp_path)
    assert created is True
    second, _rows, created_again = maybe_generate_daily(store, date(2026, 10, 21), tmp_path)
    assert created_again is False
    assert first == second


def test_daily_report_only_covers_that_day(tmp_path):
    store = Store(tmp_path / 'c.db').init()
    store.upsert({'company_name': 'Today Co', 'company_website': 'https://today.example'})
    path, rows = build_report(store, date.today(), tmp_path)
    assert rows == 1
    other, rows_other = build_report(store, date(2020, 1, 1), tmp_path)
    assert rows_other == 0


# ── pipeline ──


def test_parse_json_handles_fences_and_prose():
    assert parse_json('```json\n{"a": 1}\n```') == {'a': 1}
    assert parse_json('here you go: {"b": 2} thanks') == {'b': 2}
    assert parse_json('no json here') == {}
    assert parse_json('') == {}


def test_heuristic_queries_follow_the_rules():
    icp = ICP(industries=['SaaS', 'EdTech'], titles=['VP Sales'], geography=['United States'])
    queries = heuristic_queries(icp, 6)
    assert len(queries) >= 4
    assert all(len(q) <= 128 for q in queries)
    # Updated: no longer uses LinkedIn, should exclude job sites
    assert any('-site:indeed.com' in q for q in queries)
    assert not any('site:linkedin.com/in' in q for q in queries)  # LinkedIn removed
    assert all('intitle:' not in q for q in queries)


def test_is_directory_blocks_known_hosts():
    assert is_directory('https://www.linkedin.com/in/x')
    assert is_directory('https://crunchbase.com/organization/x')
    assert not is_directory('https://auxia.io/about')


def test_heuristic_judge_picks_the_best_non_directory():
    hits = [
        SearchHit(url='https://www.linkedin.com/in/x', authority=4, rerank_score=3.0),
        SearchHit(url='https://auxia.io', authority=2, rerank_score=1.4),
    ]
    verdict = heuristic_judge(hits)
    assert verdict['is_company'] is True
    assert verdict['company_website'] == 'https://auxia.io'
    assert verdict['company_name'] == 'Auxia'


def test_heuristic_judge_declines_when_only_directories():
    hits = [SearchHit(url='https://www.linkedin.com/in/x')]
    assert heuristic_judge(hits)['is_company'] is False


def test_extract_emails_ignores_noise():
    text = 'Mail rich@auxia.io or logo@sentry.io, see image@x.png'
    assert extract_emails(text) == ['rich@auxia.io']


def test_heuristic_extract_fills_known_fields():
    page = FetchPage(
        url='https://auxia.io',
        title='Auxia',
        content='# Auxia\n\nAuxia helps SaaS teams keep users. Contact rich@auxia.io.',
    )
    icp = ICP(industries=['SaaS'])
    record = heuristic_extract(page, icp, 'Auxia', 'q')
    assert record['company_name'] == 'Auxia'
    assert record['company_email'] == 'rich@auxia.io'
    assert record['industry'] == 'SaaS'
    assert record['match_score'] >= 55
    assert record['match_result'].startswith(rank_label(record['match_score']))


class FakeClient:
    '''Stands in for DeepSeekClient: no network, deterministic replies.'''

    def __init__(self, page_markdown, llm_replies=None):
        self.page_markdown = page_markdown
        self.llm_replies = list(llm_replies or [])
        self.prompts = []
        self.searches = []

    def check_credentials(self):
        return None

    async def search(self, query, count=10, **kwargs):
        self.searches.append(query)
        return [
            SearchHit(url='https://auxia.io', title='Auxia', authority=3, rerank_score=2.5),
            SearchHit(url='https://www.linkedin.com/in/x', title='X', authority=3),
        ]

    async def fetch(self, url):
        return FetchPage(url=url, title='Auxia', content=self.page_markdown)

    async def complete(self, prompt, **kwargs):
        self.prompts.append(prompt)
        return self.llm_replies.pop(0) if self.llm_replies else ''


def test_pipeline_run_stores_company_without_llm(tmp_path):
    store = Store(tmp_path / 'c.db').init()
    fake = FakeClient('# Auxia\n\nWe help SaaS teams. rich@auxia.io')
    pipeline = Pipeline(fake, store, ICP(industries=['SaaS']), use_llm=False)
    stats = asyncio.run(pipeline.run(2))
    assert stats['fetched'] >= 1
    assert stats['inserted'] >= 1
    rows = store.list_recent(10)
    assert rows and rows[0]['company_name']


def test_pipeline_uses_the_prompt_files_when_llm_is_on(tmp_path):
    store = Store(tmp_path / 'c.db').init()
    fake = FakeClient(
        '# Auxia\n\nWe help SaaS teams.',
        llm_replies=[
            json.dumps({'queries': ['site:linkedin.com/in "VP Sales" SaaS']}),
            json.dumps({
                'is_company': True, 'company_name': 'Auxia',
                'company_website': 'https://auxia.io', 'confidence': 90,
            }),
            json.dumps({
                'company_name': 'Auxia', 'company_website': 'https://auxia.io',
                'company_perple': 'Rich Rao（Head of Growth）',
                'company_email': 'rich@auxia.io', 'industry': 'SaaS',
                'summary': 'retention analytics', 'match_result': '高（88/100）',
                'match_score': 88,
            }),
        ],
    )
    pipeline = Pipeline(fake, store, ICP(industries=['SaaS']), use_llm=True)
    stats = asyncio.run(pipeline.run(1))
    assert fake.searches == ['site:linkedin.com/in "VP Sales" SaaS']
    assert stats['inserted'] == 1
    joined = '\n'.join(fake.prompts)
    assert '搜索提示词' in joined          # search_query.md
    assert '结果判读' in joined            # search_judge.md
    assert '公司档案提取' in joined        # fetch_extract.md
    row = store.list_recent(1)[0]
    assert row['company_perple'] == 'Rich Rao（Head of Growth）'
    assert row['match_score'] == 88


def test_pipeline_falls_back_when_llm_returns_garbage(tmp_path):
    store = Store(tmp_path / 'c.db').init()
    fake = FakeClient('# Auxia\n\nWe help SaaS teams. rich@auxia.io', ['not json', 'still no', 'nope'])
    pipeline = Pipeline(fake, store, ICP(industries=['SaaS']), use_llm=True)
    stats = asyncio.run(pipeline.run(1))
    assert stats['inserted'] == 1

def test_day_bounds_covers_rows_written_in_utc_now(tmp_path):
    from datetime import datetime, timezone

    from search_skills.report import day_bounds

    store = Store(tmp_path / 'c.db').init()
    store.upsert({'company_name': 'Just Now', 'company_website': 'https://now.example'})
    created = store.list_recent(1)[0]['created_at']
    start, end = day_bounds()
    assert start <= created < end
    path, rows = build_report(store, date.today(), tmp_path)
    assert rows == 1
